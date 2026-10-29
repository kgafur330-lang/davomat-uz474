import os, io, csv, sqlite3, secrets, json
from datetime import datetime
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

APP_DIR = os.path.dirname(os.path.abspath(__file__))
# Set DAVOMAT_DATA_DIR=/var/data on a persistent hosting disk.
DATA_DIR = os.environ.get("DAVOMAT_DATA_DIR", os.path.join(APP_DIR, "data"))
os.makedirs(DATA_DIR, exist_ok=True)
DB_PATH = os.environ.get("FACEID_DB", os.path.join(DATA_DIR, "attendance.db"))
MODEL_PATH = os.path.join(DATA_DIR, "face_model.yml")
LABELS_PATH = os.path.join(DATA_DIR, "face_labels.json")
FACES_DIR = os.path.join(DATA_DIR, "faces")
os.makedirs(FACES_DIR, exist_ok=True)
SECRET_KEY = os.environ.get("FACEID_SECRET", "CHANGE-THIS-TO-A-LONG-RANDOM-SECRET")
app = Flask(__name__)
app.secret_key = SECRET_KEY
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024

# Haar cascade is bundled with OpenCV. LBPH recognizer comes from opencv-contrib-python.
CASCADE = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
FACE_SIZE = (200, 200)
MATCH_THRESHOLD = float(os.environ.get("FACE_MATCH_THRESHOLD", "65")) # lower = stricter

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with db() as con:
        con.execute("""CREATE TABLE IF NOT EXISTS students(
            student_id TEXT PRIMARY KEY, full_name TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL)""")
        con.execute("""CREATE TABLE IF NOT EXISTS attendance(
            id INTEGER PRIMARY KEY AUTOINCREMENT, student_id TEXT NOT NULL,
            attendance_date TEXT NOT NULL, attendance_time TEXT NOT NULL,
            method TEXT NOT NULL DEFAULT 'Face ID', status TEXT NOT NULL DEFAULT 'Keldi',
            FOREIGN KEY(student_id) REFERENCES students(student_id),
            UNIQUE(student_id, attendance_date))""")
        con.execute("""CREATE TABLE IF NOT EXISTS admins(
            username TEXT PRIMARY KEY, password_hash TEXT NOT NULL)""")
        count = con.execute("SELECT COUNT(*) n FROM admins").fetchone()["n"]
        if not count:
            # Hosting should define INITIAL_ADMIN_PASSWORD as a secret environment variable.
            initial_password = os.environ.get("INITIAL_ADMIN_PASSWORD", "ChangeMe123!")
            initial_username = os.environ.get("INITIAL_ADMIN_USERNAME", "admin")
            con.execute("INSERT INTO admins(username,password_hash) VALUES(?,?)",
                        (initial_username, generate_password_hash(initial_password)))

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

def read_image(file_bytes):
    arr = np.frombuffer(file_bytes, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("Rasmni o‘qib bo‘lmadi.")
    return img

def face_crop(img):
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    faces = CASCADE.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=5, minSize=(80,80))
    if len(faces) != 1:
        raise ValueError("Rasmda aynan bitta yuz ko‘rinishi kerak. Yuzni markazga olib, yorug‘likni yaxshilang.")
    x,y,w,h = faces[0]
    crop = gray[y:y+h, x:x+w]
    crop = cv2.equalizeHist(crop)
    return cv2.resize(crop, FACE_SIZE)

def load_labels():
    if os.path.exists(LABELS_PATH):
        with open(LABELS_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def train_model():
    samples, ids, labels = [], [], {}
    with db() as con:
        students = con.execute("SELECT student_id,full_name FROM students WHERE active=1").fetchall()
    label_num = 0
    for s in students:
        folder = os.path.join(FACES_DIR, s["student_id"])
        if not os.path.isdir(folder): continue
        images = [x for x in os.listdir(folder) if x.lower().endswith((".jpg",".jpeg",".png",".webp"))]
        if not images: continue
        labels[str(label_num)] = {"student_id": s["student_id"], "full_name": s["full_name"]}
        for name in images:
            try:
                with open(os.path.join(folder,name),"rb") as f: crop = face_crop(read_image(f.read()))
                samples.append(crop); ids.append(label_num)
            except Exception:
                continue
        label_num += 1
    if not samples:
        if os.path.exists(MODEL_PATH): os.remove(MODEL_PATH)
        with open(LABELS_PATH,"w",encoding="utf-8") as f: json.dump({},f,ensure_ascii=False)
        return 0
    recognizer = cv2.face.LBPHFaceRecognizer_create()
    recognizer.train(samples, np.array(ids, dtype=np.int32))
    recognizer.save(MODEL_PATH)
    with open(LABELS_PATH,"w",encoding="utf-8") as f: json.dump(labels,f,ensure_ascii=False)
    return len(samples)

def record_attendance(student_id):
    now = datetime.now()
    try:
        with db() as con:
            con.execute("INSERT INTO attendance(student_id,attendance_date,attendance_time,method,status) VALUES(?,?,?,?,?)",
                        (student_id, now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S"), "Face ID", "Keldi"))
        return True, f"Davomat qayd etildi: {now.strftime('%Y-%m-%d %H:%M:%S')}"
    except sqlite3.IntegrityError:
        return False, "Bu talaba bugun davomatdan o‘tgan."

@app.route("/")
def index():
    with db() as con:
        total = con.execute("SELECT COUNT(*) n FROM students WHERE active=1").fetchone()["n"]
        today = con.execute("SELECT COUNT(*) n FROM attendance WHERE attendance_date=?", (datetime.now().strftime("%Y-%m-%d"),)).fetchone()["n"]
        latest = con.execute("""SELECT a.student_id,s.full_name,a.attendance_date,a.attendance_time,a.status
            FROM attendance a JOIN students s USING(student_id) ORDER BY a.id DESC LIMIT 12""").fetchall()
    return render_template("index.html", total=total, today=today, latest=latest)

@app.route("/checkin", methods=["POST"])
def checkin():
    if not request.files.get("frame"):
        return jsonify(ok=False, message="Kamera tasviri kelmadi."), 400
    if not os.path.exists(MODEL_PATH):
        return jsonify(ok=False, message="Hozircha yuz namunasi ro‘yxatga olinmagan. Administrator avval talabani ro‘yxatdan o‘tkazsin."), 400
    try:
        img = read_image(request.files["frame"].read())
        crop = face_crop(img)
        recognizer = cv2.face.LBPHFaceRecognizer_create()
        recognizer.read(MODEL_PATH)
        label_num, confidence = recognizer.predict(crop)
        labels = load_labels()
        person = labels.get(str(label_num))
        if person is None or confidence > MATCH_THRESHOLD:
            return jsonify(ok=False, message=f"Yuz mos kelmadi (score {confidence:.1f}). Qayta urinib ko‘ring."), 401
        ok, message = record_attendance(person["student_id"])
        return jsonify(ok=ok, message=message, student_id=person["student_id"], name=person["full_name"], score=round(float(confidence),1))
    except ValueError as e:
        return jsonify(ok=False, message=str(e)), 400
    except Exception as e:
        app.logger.exception("checkin error")
        return jsonify(ok=False, message="Tekshiruvda xatolik. Administratorga murojaat qiling."), 500

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username","").strip()
        password = request.form.get("password","")
        with db() as con:
            admin = con.execute("SELECT * FROM admins WHERE username=?", (username,)).fetchone()
        if admin and check_password_hash(admin["password_hash"], password):
            session["admin"] = username
            return redirect(url_for("admin"))
        flash("Login yoki parol xato.", "error")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))

@app.route("/admin", methods=["GET","POST"])
@login_required
def admin():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            sid = request.form.get("student_id","").strip().upper()
            name = request.form.get("full_name","").strip()
            if not sid or not name:
                flash("Student ID va ism-familiya kiritilishi shart.", "error")
            else:
                try:
                    with db() as con:
                        con.execute("INSERT INTO students(student_id,full_name,active,created_at) VALUES(?,?,1,?)",
                                    (sid,name,datetime.now().isoformat(timespec="seconds")))
                    flash("Talaba qo‘shildi. Endi kamida 3 ta sifatli yuz suratini yuklang.", "success")
                except sqlite3.IntegrityError:
                    flash("Bu Student ID oldin qo‘shilgan.", "error")
        elif action == "enroll":
            sid = request.form.get("student_id","").strip().upper()
            with db() as con:
                student = con.execute("SELECT * FROM students WHERE student_id=?", (sid,)).fetchone()
            if not student:
                flash("Talaba topilmadi.", "error")
            else:
                uploads = request.files.getlist("photos")
                folder = os.path.join(FACES_DIR,sid)
                os.makedirs(folder,exist_ok=True)
                saved=0
                for i,up in enumerate(uploads):
                    if not up or not up.filename: continue
                    try:
                        crop = face_crop(read_image(up.read()))
                        out = os.path.join(folder, f"enroll_{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{i}.png")
                        cv2.imwrite(out,crop); saved+=1
                    except Exception as e:
                        flash(f"{up.filename}: {e}", "error")
                if saved:
                    count=train_model()
                    flash(f"{saved} ta surat qabul qilindi. Model qayta o‘qitildi ({count} ta surat).", "success")
        elif action == "retrain":
            count=train_model()
            flash(f"Model qayta o‘qitildi. Foydalanilgan suratlar: {count}.", "success")
        elif action == "password":
            old=request.form.get("old_password",""); new=request.form.get("new_password","")
            with db() as con:
                adminrow=con.execute("SELECT * FROM admins WHERE username=?",(session["admin"],)).fetchone()
                if adminrow and check_password_hash(adminrow["password_hash"],old) and len(new)>=12:
                    con.execute("UPDATE admins SET password_hash=? WHERE username=?",(generate_password_hash(new),session["admin"]))
                    flash("Parol yangilandi.","success")
                else: flash("Eski parol xato yoki yangi parol kamida 12 belgidan iborat bo‘lishi kerak.","error")
        return redirect(url_for("admin"))
    with db() as con:
        students=con.execute("SELECT s.*, (SELECT COUNT(*) FROM attendance a WHERE a.student_id=s.student_id) attendance_count FROM students s ORDER BY s.full_name").fetchall()
    return render_template("admin.html", students=students)

@app.route("/export.xlsx")
@login_required
def export_xlsx():
    start=request.args.get("start","").strip(); end=request.args.get("end","").strip()
    query="""SELECT a.student_id,s.full_name,a.attendance_date,a.attendance_time,a.method,a.status
             FROM attendance a JOIN students s USING(student_id) WHERE 1=1"""
    args=[]
    if start: query+=" AND a.attendance_date>=?"; args.append(start)
    if end: query+=" AND a.attendance_date<=?"; args.append(end)
    query+=" ORDER BY a.attendance_date DESC,a.attendance_time DESC"
    with db() as con: rows=con.execute(query,args).fetchall()
    wb=Workbook(); ws=wb.active; ws.title="Davomat"
    ws.append(["Student ID","Talaba F.I.Sh.","Sana","Vaqt","Usul","Holat"])
    for r in rows: ws.append([r["student_id"],r["full_name"],r["attendance_date"],r["attendance_time"],r["method"],r["status"]])
    for c in ws[1]:
        c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1746A2"); c.alignment=Alignment(horizontal="center")
    for i,w in enumerate([18,38,16,14,16,16],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    output=io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name="FaceID_Davomat.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/health")
def health():
    return jsonify(status="ok", database=os.path.basename(DB_PATH))

# Initialize on import too, so production WSGI servers (Gunicorn) are ready.
init_db()

if __name__ == "__main__":
    print("DAVOMAT.UZ474: http://127.0.0.1:5000")
    print("Local default login: admin / ChangeMe123! (change it immediately)")
    app.run(host=os.environ.get("HOST","127.0.0.1"), port=int(os.environ.get("PORT","5000")), debug=False)


# ===== DAVOMAT STUDENT TELEGRAM BOT =====
import threading
import time
import urllib.request
import urllib.parse
from datetime import datetime, timezone, timedelta

UZ_TZ = timezone(timedelta(hours=5))
BOT_API = "https://api.telegram.org/bot"
student_steps = {}

def telegram_bot():
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
    admin_chat = os.environ.get("TELEGRAM_ADMIN_CHAT_ID", "").strip()

    if not token:
        print("TELEGRAM_BOT_TOKEN sozlanmagan!")
        return

    api = BOT_API + token + "/"

    def call(method, data=None):
        payload = urllib.parse.urlencode(data or {}).encode()
        req = urllib.request.Request(api + method, data=payload)
        with urllib.request.urlopen(req, timeout=35) as response:
            return json.loads(response.read().decode())

    def send(chat_id, text):
        call("sendMessage", {"chat_id": chat_id, "text": text})

    # Polling ishlashi uchun webhookni o'chirishga urinadi.
    try:
        call("deleteWebhook", {"drop_pending_updates": "false"})
    except Exception as e:
        print("Webhook sozlamasi:", e)

    with db() as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS telegram_students (
                student_id TEXT PRIMARY KEY,
                chat_id TEXT NOT NULL UNIQUE,
                registered_at TEXT NOT NULL
            )
        """)

    def normalize(text):
        return " ".join((text or "").strip().casefold().split())

    def daily_report():
        if not admin_chat:
            return
        today = datetime.now(UZ_TZ).strftime("%Y-%m-%d")
        with db() as con:
            total = con.execute(
                "SELECT COUNT(*) FROM students WHERE active=1"
            ).fetchone()[0]
            present = con.execute(
                "SELECT COUNT(*) FROM attendance WHERE attendance_date=?",
                (today,)
            ).fetchone()[0]
            rows = con.execute("""
                SELECT s.student_id, s.full_name, a.attendance_time
                FROM attendance a
                JOIN students s USING(student_id)
                WHERE a.attendance_date=?
                ORDER BY a.attendance_time
            """, (today,)).fetchall()

        lines = [
            "📊 DAVOMAT — KUNLIK HISOBOT",
            f"📅 Sana: {today}",
            f"👨‍🎓 Jami faol talabalar: {total}",
            f"✅ Kelganlar: {present}",
            f"❌ Kelmaganlar: {max(0, total-present)}",
            "",
            "KELGAN TALABALAR:"
        ]
        for row in rows[:100]:
            lines.append(
                f"{row['student_id']} | {row['full_name']} | "
                f"{row['attendance_time']}"
            )
        if not rows:
            lines.append("Bugun davomat qayd etilmagan.")

        send(admin_chat, "\n".join(lines))

    def polling():
        offset = 0
        last_report_date = ""

        while True:
            try:
                now = datetime.now(UZ_TZ)
                today = now.strftime("%Y-%m-%d")

                # Har kuni Toshkent vaqti bilan 18:00 da hisobot.
                if now.hour >= 18 and last_report_date != today:
                    daily_report()
                    last_report_date = today

                result = call("getUpdates", {
                    "offset": offset,
                    "timeout": 25
                })

                for update in result.get("result", []):
                    offset = update["update_id"] + 1
                    msg = update.get("message", {})
                    chat = msg.get("chat", {})
                    chat_id = str(chat.get("id", ""))
                    text = (msg.get("text") or "").strip()

                    if not chat_id or not text:
                        continue

                    if text == "/id":
                        send(chat_id, f"Telegram Chat ID: {chat_id}")
                        continue

                    if text.startswith("/start"):
                        student_steps[chat_id] = "id"
                        send(
                            chat_id,
                            "Assalomu alaykum! DAVOMAT botiga xush "
                            "kelibsiz.\n\nStudent ID raqamingizni yuboring."
                        )
                        continue

                    if text == "/cancel":
                        student_steps.pop(chat_id, None)
                        send(chat_id, "Amal bekor qilindi. /start bosing.")
                        continue

                    step = student_steps.get(chat_id)

                    if step == "id":
                        sid = text.upper()
                        with db() as con:
                            student = con.execute(
                                "SELECT student_id, full_name FROM students "
                                "WHERE student_id=? AND active=1",
                                (sid,)
                            ).fetchone()

                        if not student:
                            send(
                                chat_id,
                                "Bu Student ID topilmadi. ID ni tekshirib, "
                                "qaytadan yuboring."
                            )
                            continue

                        student_steps[chat_id] = {
                            "step": "name",
                            "student_id": sid
                        }
                        send(
                            chat_id,
                            "Endi ism-familiyangizni bazadagi kabi yuboring:\n"
                            + student["full_name"]
                        )
                        continue

                    if isinstance(step, dict) and step.get("step") == "name":
                        sid = step["student_id"]
                        with db() as con:
                            student = con.execute(
                                "SELECT full_name FROM students "
                                "WHERE student_id=? AND active=1",
                                (sid,)
                            ).fetchone()

                            if not student or normalize(text) != normalize(
                                student["full_name"]
                            ):
                                send(
                                    chat_id,
                                    "Ism-familiya mos kelmadi. Qaytadan "
                                    "boshlash uchun /start bosing."
                                )
                                student_steps.pop(chat_id, None)
                                continue

                            existing = con.execute(
                                "SELECT chat_id FROM telegram_students "
                                "WHERE student_id=? OR chat_id=?",
                                (sid, chat_id)
                            ).fetchone()

                            if existing and (
                                existing["chat_id"] != chat_id
                            ):
                                send(
                                    chat_id,
                                    "Bu Student ID boshqa Telegram "
                                    "akkauntiga bog'langan. Administratorga "
                                    "murojaat qiling."
                                )
                                student_steps.pop(chat_id, None)
                                continue

                            con.execute("""
                                INSERT INTO telegram_students
                                (student_id, chat_id, registered_at)
                                VALUES (?, ?, ?)
                                ON CONFLICT(student_id) DO UPDATE SET
                                chat_id=excluded.chat_id
                            """, (
                                sid, chat_id,
                                datetime.now(UZ_TZ).isoformat()
                            ))

                        student_steps.pop(chat_id, None)
                        send(
                            chat_id,
                            "✅ Ro'yxatdan o'tdingiz!\n"
                            f"Student ID: {sid}\n"
                            f"F.I.Sh.: {student['full_name']}\n\n"
                            "Davomat kamera orqali saytning o'zida "
                            "qayd etiladi. Botga /start yuborib "
                            "qayta boshlashingiz mumkin."
                        )
                        continue

                    send(chat_id, "Boshlash uchun /start bosing.")

            except Exception as e:
                print("Telegram bot xatosi:", repr(e))
                time.sleep(5)

    threading.Thread(target=polling, daemon=True).start()
    print("DAVOMAT Telegram bot ishga tushdi!")

telegram_bot()
# ===== BOT TUGADI =====
 
